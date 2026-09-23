#!/usr/bin/env python3
"""
统计 fix_D 和 fix_R 两类迁移场景中各粒度 transition 表的三阶段版本数量。

输入：
    - input/fix_D.xlsx
    - input/fix_R.xlsx

处理内容：
    1. 只处理 class/function/method 的 transition 表页；
    2. 从 H 列 JSON 的 order 中提取对应迁移场景的版本序列；
    3. 根据 Va-1 和 Vd 将版本序列划分为 before、during、after 三个阶段；
    4. 仅统计第 K 列 Key_Versions 非空的行，记录其关键版本总数及分别落在 before、during、after 阶段的数量；
    5. 输出每个用例的阶段数量以及按迁移场景、粒度汇总的阶段数量。

输出：
    output/phase_versions_summary.json
"""

import json
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List

import openpyxl


BASE = Path(__file__).resolve().parent
INPUT_FILES = {
    "fix_D": BASE / "input" / "fix_D.xlsx",
    "fix_R": BASE / "input" / "fix_R.xlsx",
}
OUTPUT_FILE = BASE / "output" / "phase_versions_summary.json"

TRANSITION_SHEETS = {
    "class_transition": "class",
    "function_transition": "function",
    "method_transition": "method",
}

# 两个 '-' 的版本字符串中，数字表示版本之间的分隔符位置（从 1 开始）。
DASH2_SPECIAL = {
    "0.16.1-0.17.1-1": 1,
    "0.17.1-1-0.20.0": 2,
    "2.0.3-1-v3.0.0": 2,
    "v0.12.1-v0.18.0-1": 1,
    "v0.13.3-v0.18.0-1": 1,
    "v0.17.1-v0.18.0-1": 1,
    "v0.18.0-1-v1.0.0": 2,
    "v0.18.0-1-v1.1.0": 2,
    "v1.0.5-2.0.3-1": 1,
}


def split_version(vstr: str) -> tuple[str, str]:
    """解析版本对字符串，返回 (左侧版本, 右侧版本)。"""
    dash_count = vstr.count("-")

    if dash_count == 1:
        left, right = vstr.split("-")
        return left, right

    if dash_count == 3:
        # 例如 networkx-2.3-networkx-3.0 -> (2.3, 3.0)
        parts = vstr.split("-")
        return parts[1], parts[3]

    if dash_count == 2:
        parts = vstr.split("-")
        split_index = DASH2_SPECIAL.get(vstr)

        if split_index is None:
            print(f"  警告: 未知 2-dash 字符串: {vstr}，使用第 1 个 '-' 分隔")
            return parts[0], "-".join(parts[1:])

        if split_index == 1:
            return parts[0], "-".join(parts[1:])
        return "-".join(parts[:2]), parts[2]

    raise ValueError(f"Unexpected dash count {dash_count}: {vstr}")


def _cell_text(value: Any) -> str:
    """将 Excel 单元格值转换为去除首尾空白的字符串。"""
    return str(value or "").strip()


def _normalize_version(version: str) -> str:
    """统一版本字符串格式，去除首尾空白及版本号前的 v。"""
    return version.strip().lstrip("v")


def _phase_for_index(
    version_idx: int,
    va1_idx: int,
    vd_idx: int,
    version_count: int,
) -> str:
    """根据版本在 Vi 序列中的位置返回 before/during/after 阶段。"""
    if va1_idx >= 0 and version_idx <= va1_idx:
        return "before"
    if vd_idx < version_count and version_idx >= vd_idx:
        return "after"
    return "during"


def count_key_versions_by_phase(
    key_versions_text: str,
    experiment: str,
    processed_versions: List[str],
    va1_idx: int,
    vd_idx: int,
) -> Dict[str, int]:
    """统计关键版本对应的 Vi 分别落在哪个阶段。"""
    counts = {"key_before": 0, "key_during": 0, "key_after": 0}

    for key_version in key_versions_text.split(","):
        key_version = key_version.strip()
        if not key_version:
            continue

        try:
            # Key_Versions 保存的是 Vi-1->Vi，只使用箭头后的 Vi。
            _, vi = key_version.split("->", 1)
            left, right = split_version(vi.strip())
            selected_version = right if experiment == "fix_D" else left
            selected_version = _normalize_version(selected_version)
            version_idx = processed_versions.index(selected_version)
        except (TypeError, ValueError):
            # 保持原有行为：无法解析或无法在 order 中定位的关键版本不参与阶段统计。
            continue

        phase = _phase_for_index(
            version_idx=version_idx,
            va1_idx=va1_idx,
            vd_idx=vd_idx,
            version_count=len(processed_versions),
        )
        counts[f"key_{phase}"] += 1

    return counts


def process_sheet(
    ws: Any,
    experiment: str,
    granularity: str,
) -> List[Dict[str, Any]]:
    """处理一个 transition 表页，返回每个有效用例的阶段数量。"""
    is_fix_d = experiment == "fix_D"
    results: List[Dict[str, Any]] = []

    for row_idx in range(2, ws.max_row + 1):
        fqn_d = _cell_text(ws.cell(row_idx, 1).value)
        fqn_r = _cell_text(ws.cell(row_idx, 2).value)
        va1 = _cell_text(ws.cell(row_idx, 5).value)
        vd = _cell_text(ws.cell(row_idx, 6).value)
        ranking_json = ws.cell(row_idx, 8).value
        key_versions_text = _cell_text(ws.cell(row_idx, 11).value)

        if not fqn_d or not fqn_r or not va1 or not vd or not ranking_json or not key_versions_text:
            continue

        try:
            data = json.loads(ranking_json)
        except (json.JSONDecodeError, TypeError):
            continue

        order = data.get("order", [])
        if not isinstance(order, list) or not order:
            continue

        processed_versions: List[str] = []
        try:
            for version_pair in order:
                left, right = split_version(str(version_pair))
                version = right if is_fix_d else left
                processed_versions.append(_normalize_version(version))
        except (TypeError, ValueError):
            continue

        try:
            va1_idx = processed_versions.index(va1)
        except ValueError:
            va1_idx = -1

        try:
            vd_idx = processed_versions.index(vd)
        except ValueError:
            vd_idx = len(processed_versions)

        before = va1_idx + 1 if va1_idx >= 0 else 0
        after = len(processed_versions) - vd_idx if vd_idx < len(processed_versions) else 0
        during = len(processed_versions) - before - after

        key_version_count = len([item for item in key_versions_text.split(",") if item.strip()])
        key_phase_counts = count_key_versions_by_phase(
            key_versions_text=key_versions_text,
            experiment=experiment,
            processed_versions=processed_versions,
            va1_idx=va1_idx,
            vd_idx=vd_idx,
        )

        results.append(
            {
                "case": f"{fqn_d}-{fqn_r}",
                "column_1": fqn_d,
                "column_2": fqn_r,
                "granularity": granularity,
                "sheet": ws.title,
                "row": row_idx,
                "before": before,
                "during": during,
                "after": after,
                "key_version_count": key_version_count,
                **key_phase_counts,
            }
        )

    return results


def empty_counts() -> Dict[str, int]:
    return {
        "case_count": 0,
        "before": 0,
        "during": 0,
        "after": 0,
        "key_version_count": 0,
        "key_before": 0,
        "key_during": 0,
        "key_after": 0,
    }


def add_counts(target: Dict[str, int], source: Dict[str, Any]) -> None:
    target["case_count"] += 1
    for field in (
        "before",
        "during",
        "after",
        "key_version_count",
        "key_before",
        "key_during",
        "key_after",
    ):
        target[field] += int(source[field])


def build_group(experiment: str, granularity: str, entries: List[Dict[str, Any]]) -> Dict[str, Any]:
    """构造一个迁移场景-粒度分组，包含用例明细和该组汇总。"""
    case_map: Dict[str, Dict[str, Any]] = {}
    for entry in entries:
        case_id = entry["case"]
        if case_id in case_map:
            print(
                f"  警告: {experiment}/{granularity} 中发现重复用例 {case_id}，保留首次记录"
            )
            continue
        case_map[case_id] = entry

    cases = list(case_map.values())
    summary = empty_counts()
    for case in cases:
        add_counts(summary, case)

    return {
        "cases": cases,
        "summary": summary,
    }


def empty_experiment() -> Dict[str, Any]:
    return {
        "class": {"cases": [], "summary": empty_counts()},
        "function": {"cases": [], "summary": empty_counts()},
        "method": {"cases": [], "summary": empty_counts()},
    }


def process_experiment(experiment: str, input_path: Path) -> Dict[str, Any]:
    """读取一个 fix_D/fix_R 文件，只处理 transition 表页。"""
    if not input_path.is_file():
        raise FileNotFoundError(f"输入文件不存在: {input_path}")

    grouped: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    workbook = openpyxl.load_workbook(input_path, data_only=True)

    try:
        for sheet_name, granularity in TRANSITION_SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                continue
            rows = process_sheet(
                ws=workbook[sheet_name],
                experiment=experiment,
                granularity=granularity,
            )
            grouped[granularity].extend(rows)
    finally:
        workbook.close()

    return {
        granularity: build_group(experiment, granularity, grouped[granularity])
        for granularity in ("class", "function", "method")
    }


def build_aggregate_summary(report: Dict[str, Any]) -> Dict[str, Any]:
    """构造最终汇总数据，不区分 same/diff 或 no_transition。"""
    aggregate: Dict[str, Any] = {}

    for experiment in ("fix_D", "fix_R"):
        experiment_summary: Dict[str, Any] = {}
        overall = empty_counts()

        for granularity in ("class", "function", "method"):
            summary = report[experiment][granularity]["summary"]
            experiment_summary[granularity] = summary.copy()
            for field in (
                "before",
                "during",
                "after",
                "key_version_count",
                "key_before",
                "key_during",
                "key_after",
            ):
                overall[field] += summary[field]
            overall["case_count"] += summary["case_count"]

        experiment_summary["overall"] = overall
        aggregate[experiment] = experiment_summary

    return aggregate


def main() -> None:
    report: Dict[str, Any] = {
        "fix_D": process_experiment("fix_D", INPUT_FILES["fix_D"]),
        "fix_R": process_experiment("fix_R", INPUT_FILES["fix_R"]),
    }

    # 将原脚本的汇总信息放在 JSON 的最后，便于直接查看总体统计。
    report["summary"] = build_aggregate_summary(report)

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_FILE.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"完成，结果已保存到: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
