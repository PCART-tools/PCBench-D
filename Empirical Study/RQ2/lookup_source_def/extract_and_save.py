import argparse
import logging
from pathlib import Path
from typing import Optional

import openpyxl

# 引入提取工具模块
from tool.extract_dep import get_api_source_by_url_fqn_version_type
from tool.extract_and_save_rep import extract_candidates
from trim_large_candidates import clean_large_candidates

# 日志配置将在 main 函数中动态初始化


def compute_prefix(fqn_d: str, fqn_r: str, api_type: int) -> str:
    """
    功能：
        计算用于本地仓库扫描的相对路径前缀 (prefix)。
        
    逻辑：
        1. 将 fqn_d 和 fqn_r 按 '.' 分隔为列表。
        2. 比较两者段数，选取段数较多的列表。
        3. 根据 api_type 去除末尾段数 (api_type=3去3段，其余去2段)。
        4. 剩余部分用 '/' 连接。如果截断后为空，则返回 "."。

    参数：
        fqn_d (str): 废弃 API 的全限定名。
        fqn_r (str): 候选 API 的全限定名。
        api_type (int): API 类型 (1=函数, 2=类, 3=方法)。

    返回：
        str: 拼接好的路径前缀字符串。
    """
    fqn_d_parts = fqn_d.split('.') if fqn_d else []
    fqn_r_parts = fqn_r.split('.') if fqn_r else []
    
    # 选取段数较多的
    target_parts = fqn_d_parts if len(fqn_d_parts) >= len(fqn_r_parts) else fqn_r_parts
    
    # 去除段数
    drop_count = 3 if api_type == 3 else 2
    prefix_parts = target_parts[:-drop_count] if len(target_parts) > drop_count else []
    
    return "/".join(prefix_parts) if prefix_parts else "."


def parse_versions(version_str: Optional[str]) -> list:
    """
    功能：
        解析按逗号分隔的版本字符串，返回保持原有顺序的有效版本列表。
    """
    if not version_str:
        return []
    return [v.strip() for v in str(version_str).split(',') if v.strip()]


def process_excel(excel_path: Path, api_type: int, output_dir: Path, repo_root: Path):
    """
    功能：
        解析 Excel 文件，并逐行调用远程及本地提取工具，保存 API 源码。

    参数：
        excel_path (Path): 输入的 Excel 文件路径。
        api_type (int): 提取的 API 粒度 (1=函数, 2=类, 3=方法)。
        output_dir (Path): 提取结果的存放根目录。
        repo_root (Path): 本地 Git 仓库的根目录。
    """
    logging.info(f"正在加载 Excel 文件: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        libname = sheet_name
        logging.info(f"======== 开始处理表页 (库名): {libname} ========")
        
        # 遍历每一行 (跳过第一行表头，索引从2开始)
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            # 判断空行
            if not row or not row[0]:
                continue
            
            try:
                # 提取列数据 (注意：Python 索引从 0 开始)
                # Excel 对应: 1:fqn_d(0), 2:fqn_r(1), 7:url(6), 9:v_bef(8), 10:v_trans(9), 11:v_aft(10)
                fqn_d = str(row[0]).strip() if row[0] else ""
                fqn_r = str(row[1]).strip() if len(row) > 1 and row[1] else ""
                url = str(row[6]).strip() if len(row) > 6 and row[6] else ""
                
                v_bef = str(row[8]).strip() if len(row) > 8 and row[8] else ""
                v_trans = str(row[9]).strip() if len(row) > 9 and row[9] else ""
                v_aft = str(row[10]).strip() if len(row) > 10 and row[10] else ""
                
                if not fqn_d or not fqn_r:
                    logging.warning(f"第 {row_idx} 行缺少 fqn_d 或 fqn_r，跳过。")
                    continue
                
                # 2.1 创建基础输出目录结构
                # output/libname/fqn_d-fqn_r/
                base_path = output_dir / libname / f"{fqn_d}-{fqn_r}"
                source_path = base_path / fqn_d
                cans_path = base_path / "R_candidates"
                
                source_path.mkdir(parents=True, exist_ok=True)
                cans_path.mkdir(parents=True, exist_ok=True)
                
                # 2.2 提取废弃 API 源码 (bef_n_versions + transition_versions，保持顺序 bef < trans)
                dep_versions = parse_versions(v_bef) + parse_versions(v_trans)
                for version in dep_versions:
                    out_file = source_path / f"{version}.py"
                    if out_file.exists():
                        logging.debug(f"已存在废弃 API 文件 {out_file}，跳过。")
                        continue
                    
                    if not url:
                        logging.warning(f"第 {row_idx} 行 url 为空，无法提取 {version} 的废弃 API。")
                        continue
                    
                    try:
                        logging.info(f"远程提取废弃 API: {libname} | {fqn_d} | {version}")
                        source_code = get_api_source_by_url_fqn_version_type(url, fqn_d, version, api_type)
                        out_file.write_text(source_code, encoding="utf-8")
                    except Exception as e:
                        logging.error(f"提取废弃 API 失败 ({libname}, {fqn_d}, {version}): {e}")
                
                # 2.3 提取候选 API 源码 (transition_versions + aft_n_versions，保持顺序 trans < aft)
                rep_versions = parse_versions(v_trans) + parse_versions(v_aft)
                prefix = compute_prefix(fqn_d, fqn_r, api_type)
                
                for version in rep_versions:
                    cur_cans_path = cans_path / version
                    # 判断目录是否存在且非空
                    if cur_cans_path.exists() and any(cur_cans_path.iterdir()):
                        logging.debug(f"已存在候选 API 目录且非空 {cur_cans_path}，跳过。")
                        continue
                    
                    cur_cans_path.mkdir(parents=True, exist_ok=True)
                    
                    try:
                        logging.info(f"本地批量提取候选 API: {libname} | prefix: {prefix} | {version}")
                        # 调用本地批量提取
                        extract_candidates(
                            root=repo_root,
                            libname=libname,
                            version=version,
                            path=prefix,
                            type=api_type,
                            outputRoot=cur_cans_path
                        )
                    except Exception as e:
                        logging.error(f"提取候选 API 失败 ({libname}, {version}, prefix={prefix}): {e}")

            except Exception as e:
                logging.error(f"处理表页 {libname} 第 {row_idx} 行时发生未捕获异常: {e}")


def main():
    parser = argparse.ArgumentParser(description="从 Excel 读取并提取 API 的废弃与候选源码。")
    parser.add_argument(
        "--excel", 
        type=str, 
        required=True, 
        help="输入的 Excel 文件路径"
    )
    parser.add_argument(
        "--api_type", 
        type=int, 
        choices=[1, 2, 3], 
        required=True, 
        help="提取的 API 粒度：1=函数, 2=类, 3=方法"
    )
    parser.add_argument(
        "--output_dir", 
        type=str, 
        default="/media/he/Rbench/similarity/output", 
        help="源码输出的根目录，默认为 /media/he/Rbench/similarity/output"
    )
    parser.add_argument(
        "--repo_root", 
        type=str, 
        default="/dataset/he", 
        help="存放本地仓库代码的根目录，默认为 /dataset/he"
    )
    
    args = parser.parse_args()
    
    excel_path = Path(args.excel).resolve()
    
    # 动态设置日志文件名为传入的 Excel 表名，保存到当前脚本同级目录
    log_file_path = Path(__file__).parent / f"{excel_path.stem}.log"
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s',
        handlers=[
            logging.FileHandler(log_file_path, encoding='utf-8'),
            logging.StreamHandler()
        ],
        force=True  # 确保在已有其他模块配置 logging 的情况下强制覆盖
    )
    
    # 根据 api_type 动态调整 output 目录
    base_output_dir = Path(args.output_dir).resolve()
    type_map = {1: "function", 2: "class", 3: "method"}
    output_dir = base_output_dir / type_map[args.api_type]
    
    # 询问子路径
    print(f"当前 output 路径为: {output_dir}")
    sub_path = input("请输入子路径 (直接回车将使用自动编号): ").strip()
    
    if sub_path:
        output_dir = output_dir / sub_path
    else:
        # 如果未输入，计算已有目录数量作为编号
        output_dir.mkdir(parents=True, exist_ok=True)
        count = sum(1 for item in output_dir.iterdir() if item.is_dir())
        output_dir = output_dir / str(count)
    
    repo_root = Path(args.repo_root).resolve()
    
    if not excel_path.exists():
        logging.error(f"指定的 Excel 文件不存在: {excel_path}")
        return
        
    logging.info(f"Excel 路径: {excel_path}")
    logging.info(f"API 类型: {args.api_type}")
    logging.info(f"输出根目录: {output_dir}")
    logging.info(f"本地仓库根目录: {repo_root}")
    
    # 执行提取处理
    process_excel(excel_path, args.api_type, output_dir, repo_root)
    logging.info("所有表页处理完成。")

    # 询问是否执行 trim_large_candidates
    trim_ans = input(f"是否对 {output_dir} 执行大文件清理 (trim_large_candidates)? [y/N]: ").strip().lower()
    if trim_ans in ('y', 'yes'):
        logging.info("开始执行 trim_large_candidates...")
        try:
            clean_large_candidates(output_dir)
            logging.info("大文件清理执行完成。")
        except Exception as e:
            logging.error(f"执行大文件清理失败: {e}")
    else:
        logging.info("跳过大文件清理。")


if __name__ == "__main__":
    main()
