import argparse
import os

import openpyxl

# ===================== 路径配置（可用命令行参数覆盖） =====================
#   python fqn_classify.py --input <输入目录> --output <输出目录>
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_INPUT_DIR = os.path.join(BASE_DIR, 'input')
DEFAULT_OUTPUT_DIR = os.path.join(BASE_DIR, 'output')

# 输入文件名 -> 粒度
FILE_GRANULARITY = {
    'class.xlsx': 'class',
    'function.xlsx': 'function',
    'method.xlsx': 'method',
}

# 需要跳过的表页（不参与分类）
SKIP_SHEETS = {'summary'}

# type 列的候选表头：已有则复用；没有（input 三个文件均已去掉该列）则在末尾新增一列
TYPE_HEADERS = {'result', 'types'}
NEW_TYPE_HEADER = 'types'


def _classify_class_or_function(segs1, segs2):
    """
    class 和 function 的 type 含义完全相同：
      type0: identical（完全相同）
      type1: same name, different module（同名，不同模块）
      type2: different name, same module（同模块，不同名）
      type3: different name and module（不同名，不同模块）

    分段比较规则（将 FQN 用 '.' 拆分为若干分段，最后一个分段是名字，前面是模块路径）：
      type0 = 所有分段完全相同；
      type1 = 最后一个分段（名字）相同，前面分段（模块路径）有所不同；
      type2 = 除了最后一个分段外，前面分段完全相同（即仅名字不同，模块相同）；
      type3 = 最后一个分段不同，且前面分段也存在不同。
    """
    if segs1 == segs2:
        return 0
    if segs1[-1] == segs2[-1]:
        # 最后一个分段（名字）相同；且前面分段有所不同（否则上面 type0 已返回）
        return 1
    if segs1[:-1] == segs2[:-1]:
        # 最后一个分段（名字）不同，但前面分段（模块路径）完全相同
        return 2
    # 最后一个分段不同，前面分段也不同
    return 3


def _classify_method(segs1, segs2):
    """
    method 的 type 含义（FQN 形如 module[.module...].class.method，
    最后一个分段是方法名，倒数第二个分段是类名，之前是模块路径）：
      type1: same module and class, different method name（同模块同类，不同方法名）
      type2: same module, different class, same method name（同模块，不同类，同方法名）
      type3: same module, different class and method name（同模块，不同类不同方法名）
      type4: different module, same method name（不同模块，同方法名）
      type5: different module and method name（不同模块，不同方法名）
      注意：method 的两个 FQN 完全相同（同模块同类同方法名）属于错误，
      不是有效 type，直接抛出 ValueError。

    分段比较规则（将 FQN 用 '.' 拆分为若干分段）：
      type1 = 除了最后一个分段（方法名）不同，其余分段完全相同；
      type2 = 除了倒数第二个分段（类名）不同，其余分段完全相同；
      type3 = 除了最后两个分段（类名与方法名）不同，其余分段完全相同；
      type4 = 不属于 type1-3，且最后一个分段相同；
      type5 = 不属于 type1-3，且最后一个分段不相同。
    """
    if segs1 == segs2:
        # 完全相同：同模块、同类、同方法名。对 method 来说这属于错误情况，
        # 两个待比较的 method 不应该相同，抛异常提示而不是返回某个 type。
        raise ValueError(f"method 的两个 FQN 完全相同（不是有效类型）：{'.'.join(segs1)}")
    if segs1[:-1] == segs2[:-1]:
        # 仅最后一个分段（方法名）不同，其余完全相同
        return 1
    if segs1[:-2] == segs2[:-2]:
        # 到这里倒数第二个分段（类名）必然不同，
        # 否则 segs1[:-1] == segs2[:-1] 会被上面的 type1 分支捕获。
        if segs1[-1] == segs2[-1]:
            # 仅类名不同，方法名相同
            return 2
        # 类名与方法名都不同
        return 3
    # 不属于 type1-3
    if segs1[-1] == segs2[-1]:
        # 最后一个分段（方法名）相同
        return 4
    # 最后一个分段（方法名）不同
    return 5


def classify_fqn_type(fqn1: str, fqn2: str, granularity: str) -> int:
    """
    根据两个 FQN（全限定名）和粒度，分类它们之间的关系类型。

    参数：
        fqn1, fqn2 : 全限定名，如 'com.example.Foo' / 'com.example.Bar.baz'
        granularity: 'class' 或 'function' 或 'method'

    返回：
        class/function -> type 0-3；method -> type 1-5（含义见各函数注释）。
        注意：method 若两个 FQN 完全相同，属于错误，抛出 ValueError。
    """
    # 去掉首尾空白与首尾的 '.'（如 'xxx.method.' 的尾点），再按 '.' 拆段
    segs1 = fqn1.strip().strip('.').split('.')
    segs2 = fqn2.strip().strip('.').split('.')

    if granularity in ('class', 'function'):
        return _classify_class_or_function(segs1, segs2)
    if granularity == 'method':
        return _classify_method(segs1, segs2)
    raise ValueError(f"未知粒度: {granularity!r}，应为 'class' / 'function' / 'method'")


# ===================== 批量处理（input -> output） =====================

def _find_or_create_type_col(ws):
    """返回 type 列号：若最后一列表头是 result/types 则复用，否则在末尾新增一列。"""
    last = ws.max_column
    if ws.cell(row=1, column=last).value in TYPE_HEADERS:
        return last
    col = last + 1
    ws.cell(row=1, column=col, value=NEW_TYPE_HEADER)
    return col


def _classify_row(fqn1, fqn2, granularity):
    """对一行第 1、2 列调用分类函数，返回格式化后的 type（如 'type1'）；
    FQN 为空返回 None；method 完全相同等异常返回错误文本。"""
    if fqn1 is None or fqn2 is None or str(fqn1).strip() == '' or str(fqn2).strip() == '':
        return None
    try:
        return 'type%d' % classify_fqn_type(str(fqn1), str(fqn2), granularity)
    except ValueError as e:
        return 'ERROR: %s' % e


def process_file(input_path, output_path, granularity):
    """读取 input 文件，逐表页逐行分类并填入最后一列，写入 output 文件。返回处理行数。"""
    wb = openpyxl.load_workbook(input_path, data_only=False)
    n_rows = 0
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        col = _find_or_create_type_col(ws)
        for r in range(2, ws.max_row + 1):
            v = _classify_row(ws.cell(row=r, column=1).value,
                              ws.cell(row=r, column=2).value, granularity)
            if v is not None:
                ws.cell(row=r, column=col, value=v)
                n_rows += 1
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    wb.close()
    return n_rows


def main():
    parser = argparse.ArgumentParser(
        description='读取 input 目录中的 xlsx，逐表页逐行计算 FQN 关系类型并填入最后一列，'
                    '结果写入 output 目录（input -> output）。')
    parser.add_argument('--input', default=DEFAULT_INPUT_DIR, help='输入目录（默认: input/）')
    parser.add_argument('--output', default=DEFAULT_OUTPUT_DIR, help='输出目录（默认: output/）')
    args = parser.parse_args()

    print(f'input  : {args.input}')
    print(f'output : {args.output}')

    for filename, granularity in FILE_GRANULARITY.items():
        in_path = os.path.join(args.input, filename)
        out_path = os.path.join(args.output, filename)
        if not os.path.exists(in_path):
            print(f'\n[跳过] {filename}: input 中不存在')
            continue
        n_rows = process_file(in_path, out_path, granularity)
        print(f'{filename} (granularity={granularity}): 处理 {n_rows} 行 -> {out_path}')

    print('\n完成 ✓')


if __name__ == '__main__':
    main()
