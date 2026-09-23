import argparse
import importlib.util
import sys
from pathlib import Path

from openpyxl import load_workbook


def normalize_signature(sig: str) -> str:
    """
    功能说明：
        格式化提取出的签名字符串，移除回车/换行/制表符等无意义字符，避免影响后续解析。

    参数说明：
        sig (str):
            形如 "(a: int, b=1)" 的签名字符串。

    返回：
        str:
            清洗后的签名字符串。
    """
    return (
        sig.replace("\r", "")
        .replace("\n", "")
        .replace("\t", "")
        .replace("\f", "")
        .replace("\v", "")
        .replace(" ", "")
    )


def extract_outer_parens(text: str) -> str | None:
    """
    功能说明：
        从输入字符串中提取最外层的一对圆括号及其包裹内容，例如从
        "foo(a: int, b=1) -> bool" 中提取 "(a: int, b=1)"。

    参数说明：
        text (str):
            输入字符串，通常来自 Excel 单元格内容。

    返回：
        str | None:
            若能找到匹配的最外层括号对则返回包含括号的子串；否则返回 None。
    """
    start = text.find("(")
    if start < 0:
        return None

    depth = 0
    for idx in range(start, len(text)):
        ch = text[idx]
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
            if depth == 0:
                return text[start : idx + 1]

    return None


def load_change_analyzer():
    """
    功能说明：
        通过文件路径动态加载 changeAnalyze.py 并返回其 findDiffer 函数。

    参数说明：
        无。

    返回：
        callable:
            changeAnalyze.findDiffer(oldPara: str, newPara: str) -> dict。

    异常：
        FileNotFoundError:
            changeAnalyze.py 不存在。
        ImportError:
            模块加载失败或缺少 findDiffer。
    """
    module_path = Path("/dataset/he/Rbench/R1/PCART/Change/changeAnalyze.py")
    if not module_path.exists():
        raise FileNotFoundError(str(module_path))

    spec = importlib.util.spec_from_file_location("pcart_changeAnalyze", str(module_path))
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载模块: {module_path}")

    pcart_root = str(module_path.parent.parent)
    inserted = False
    try:
        if pcart_root not in sys.path:
            sys.path.insert(0, pcart_root)
            inserted = True

        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
    finally:
        if inserted:
            if sys.path and sys.path[0] == pcart_root:
                sys.path.pop(0)
            else:
                try:
                    sys.path.remove(pcart_root)
                except ValueError:
                    pass

    find_differ = getattr(module, "findDiffer", None)
    if find_differ is None:
        raise ImportError(f"模块缺少 findDiffer: {module_path}")

    return find_differ


def analyze_workbook(input_xlsx: str | Path, out_dir: str | Path) -> Path:
    """
    功能说明：
        逐表页逐行处理 Excel：从第 6/7 列文本中提取最外层括号内容作为
        dep_signature/rep_signature，调用 changeAnalyze.findDiffer 并将返回
        dict 的字符串形式写入第 9 列，最后保存到输出目录。

    参数说明：
        input_xlsx (str | Path):
            输入 Excel 路径（多表页，首行为表头）。
        out_dir (str | Path):
            输出目录路径。若不存在会自动创建。

    返回：
        Path:
            输出 Excel 的完整路径。
    """
    input_path = Path(input_xlsx)
    out_path = Path(out_dir)
    out_path.mkdir(parents=True, exist_ok=True)

    find_differ = load_change_analyzer()

    wb = load_workbook(input_path)
    for ws in wb.worksheets:
        for row_idx in range(2, ws.max_row + 1):
            cell6 = ws.cell(row=row_idx, column=6)
            cell7 = ws.cell(row=row_idx, column=7)
            result_cell = ws.cell(row=row_idx, column=9)

            raw6 = cell6.value
            raw7 = cell7.value

            if not isinstance(raw6, str) or not isinstance(raw7, str):
                result_cell.value = repr({"error": "cell_value_not_str"})
                continue

            dep_signature = extract_outer_parens(raw6)
            rep_signature = extract_outer_parens(raw7)
            if dep_signature is None or rep_signature is None:
                result_cell.value = repr(
                    {
                        "error": "signature_not_found",
                        "dep_signature": dep_signature,
                        "rep_signature": rep_signature,
                    }
                )
                continue
            dep_signature = normalize_signature(dep_signature)
            rep_signature = normalize_signature(rep_signature)

            try:
                diff = find_differ(dep_signature, rep_signature)
                result_cell.value = repr(diff)
            except Exception as exc:
                result_cell.value = repr({"error": "analyze_failed", "message": str(exc)})

    output_file = out_path / f"{input_path.stem}_analyzed{input_path.suffix}"
    wb.save(output_file)
    return output_file


def main() -> None:
    """
    功能说明：
        脚本入口：读取输入 Excel 与输出目录，执行逐行变更分析并写回第 9 列。

    参数说明：
        --input (str):
            输入 Excel 路径。
        --out_dir (str):
            输出目录路径。

    返回：
        None:
            脚本会将输出 Excel 写入 out_dir，并在 stdout 打印输出路径。
    """
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="输入 Excel 路径")
    parser.add_argument("--out_dir", required=True, help="输出目录路径")
    args = parser.parse_args()

    output_file = analyze_workbook(args.input, args.out_dir)
    print(str(output_file))


if __name__ == "__main__":
    main()
